#%%
import json
from docx import Document



#%%
client_filename = "../data/assessments/clients/2021-richner-will.docx"
client_filename = "../data/assessments/clients/2022-kinkel-harrison.docx"
client_filename = "../data/assessments/clients/2023-manion-coleton.docx"
client_filename = "../data/assessments/clients/2024-plese-bentley.docx"
client_filename = "../data/assessments/clients/2024-hilton-brooklyn.docx"
document = Document(client_filename)

doc_tables = []
for idx, table in enumerate(document.tables):
    keys = []
    data = []
    num_rows = len(table.rows)
    num_cols = len(table.columns)
    print(idx, num_rows, num_cols)

    for r in range(num_rows):
        row_data = []
        for c in range(num_cols):
            text = table.cell(r,c).text.strip()
            text = text.replace('\t', "")
            text = text.replace('\n', "")
            text = text.replace('\u2019', "'")
            text = text.replace('\u2013', "-")

            if r == 0: # the first row contains headers
                keys.append(text)
            else:
                row_data.append(text)
        
        if len(row_data) == len(keys):
            data.append(dict(zip(keys, row_data)))

    table_info = {'rows': num_rows-1, 'cols': keys, 'data': data}
    doc_tables.append(table_info)

json_filename = "../data/client_tables.json"
with open(json_filename, 'w', encoding='utf-8') as f:
    json.dump(doc_tables, f, ensure_ascii=False, indent=2)


#%%
